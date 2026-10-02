"""
src/ingestion/parser.py

Purpose:
    Multi-format transcript file parser reading raw text from .txt, .md, .docx, and .pdf files.

Working & Flow:
    - `parse_file(file_path)` inspects file extension.
    - Routes reading logic to format-specific handler (_parse_txt, _parse_docx, _parse_pdf).
    - Returns raw text content string.

Links to:
    - src/preprocessing/cleaner.py
    - src/services/extraction_service.py
"""

import os
from pathlib import Path
import pypdf
import docx


class IngestionParser:
    """Parser class supporting txt, md, docx, and pdf file extraction."""

    @staticmethod
    def parse_file(file_path: str | Path, content_bytes: bytes | None = None) -> str:
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext in [".txt", ".md"]:
            return IngestionParser._parse_txt(file_path, content_bytes)
        elif ext == ".docx":
            return IngestionParser._parse_docx(file_path, content_bytes)
        elif ext == ".pdf":
            return IngestionParser._parse_pdf(file_path, content_bytes)
        else:
            raise ValueError(f"Unsupported file format: '{ext}'. Supported: .txt, .md, .docx, .pdf")

    @staticmethod
    def _parse_txt(file_path: str | Path, content_bytes: bytes | None) -> str:
        if content_bytes:
            return content_bytes.decode("utf-8", errors="replace")
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    @staticmethod
    def _parse_docx(file_path: str | Path, content_bytes: bytes | None) -> str:
        import io
        if content_bytes:
            doc = docx.Document(io.BytesIO(content_bytes))
        else:
            doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

    @staticmethod
    def _parse_pdf(file_path: str | Path, content_bytes: bytes | None) -> str:
        import io
        stream = io.BytesIO(content_bytes) if content_bytes else open(file_path, "rb")
        try:
            reader = pypdf.PdfReader(stream)
            text_pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(text_pages)
        finally:
            if not content_bytes:
                stream.close()
